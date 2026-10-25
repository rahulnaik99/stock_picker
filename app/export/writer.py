"""CSV (main ranking table) and Excel workbook (Rankings, Agent Scores, Events, Data Quality, Configuration)."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.ledger.events import EventLedger
from app.schemas import AGENT_NAMES, StockResult
from app.scoring import policy
from app.service.settings import get_settings

RANK_COLS = ["rank", "rank_change", "symbol", "company_name", "sub_sector", "analysis_timestamp", "fundamental_score", "industry_score",
             "indicator_score", "investment_trend_score", "option_chain_score", "final_score", "llm_proposed_score", "direction",
             "consensus_status", "confidence", "trade_status", "reasoning_summary", "risk_flags", "next_catalyst", "data_completeness",
             "data_status", "source_freshness", "run_id"]


def ranking_frame(stocks: list[StockResult], previous: Path | None = None) -> pd.DataFrame:
    prev: dict[str, int] = {}
    if previous and previous.exists():
        old = pd.read_csv(previous)
        prev = {r.symbol: int(r["rank"]) for _, r in old.dropna(subset=["rank"]).iterrows()} if "rank" in old else {}
    rows, n = [], 0
    for s in stocks:
        if s.valid:
            n += 1
        rows.append({
            "rank": n if s.valid else None,
            "rank_change": (prev[s.symbol] - n) if s.valid and s.symbol in prev else None,  # + means moved up
            "symbol": s.symbol, "company_name": s.company_name, "sub_sector": s.sub_sector, "analysis_timestamp": s.analysis_timestamp,
            **{f"{k}_score": s.scores.get(k) for k in AGENT_NAMES},
            "final_score": s.final_score, "llm_proposed_score": s.llm_proposed_score, "direction": s.direction,
            "consensus_status": s.consensus_status, "confidence": s.confidence, "trade_status": s.trade_status,
            "reasoning_summary": s.reasoning_summary, "risk_flags": " | ".join(s.risk_flags), "next_catalyst": s.next_catalyst,
            "data_completeness": s.data_completeness, "data_status": s.data_status, "source_freshness": s.source_freshness, "run_id": s.run_id})
    return pd.DataFrame(rows, columns=RANK_COLS)


def agent_scores_frame(stocks: list[StockResult]) -> pd.DataFrame:
    rows = []
    for s in stocks:
        for name in AGENT_NAMES:
            r = s.agent_results.get(name)
            rows.append({"symbol": s.symbol, "agent": name, "status": r.status if r else "missing", "score": r.score if r else None,
                         "confidence": r.confidence if r else None, "reasoning": r.reasoning if r else "", "model": f"{r.provider}/{r.model}" if r and r.provider else "",
                         "missing_data": "; ".join(r.missing_data) if r else "", "evidence": json.dumps(r.evidence, default=str)[:1500] if r else "",
                         "note": r.note if r else ""})
    return pd.DataFrame(rows)


def data_quality_frame(stocks: list[StockResult]) -> pd.DataFrame:
    rows = []
    for s in stocks:
        for name, r in s.agent_results.items():
            for src in r.sources:
                if src.get("status") != "ok":
                    rows.append({"symbol": s.symbol, "agent": name, "tool": src.get("tool"), "status": src.get("status"),
                                 "source": src.get("source"), "message": src.get("message"), "retrieved_at": src.get("retrieved_at")})
        if s.error:
            rows.append({"symbol": s.symbol, "agent": "-", "tool": "-", "status": "error", "source": "-", "message": s.error, "retrieved_at": ""})
        if s.status_adjustments:
            rows.append({"symbol": s.symbol, "agent": "policy", "tool": "-", "status": "adjusted", "source": "-", "message": " | ".join(s.status_adjustments), "retrieved_at": ""})
    return pd.DataFrame(rows, columns=["symbol", "agent", "tool", "status", "source", "message", "retrieved_at"])


def export(stocks: list[StockResult], run_id: str, out_dir: Path | None = None) -> dict[str, Path]:
    s = get_settings()
    out_dir = out_dir or s.path(s.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path, xlsx_path, latest = out_dir / f"ranking_{run_id}.csv", out_dir / f"stock_scores_{run_id}.xlsx", out_dir / "latest_ranking.csv"
    rank_df = ranking_frame(stocks, previous=latest)
    rank_df.to_csv(csv_path, index=False)
    cfg = policy.load_config()
    symbols = [x.symbol for x in stocks]
    as_of = stocks[0].analysis_timestamp if stocks else ""
    events = EventLedger().all_for(symbols)
    config_df = pd.DataFrame([{"key": "run_id", "value": run_id}, {"key": "analysis_timestamp", "value": as_of},
                              {"key": "weights_version", "value": cfg["version"]},
                              *[{"key": f"weight.{k}", "value": v} for k, v in cfg["weights"].items()],
                              {"key": "required_agents", "value": ", ".join(cfg["required"])},
                              {"key": "min_data_completeness", "value": cfg["min_data_completeness"]},
                              {"key": "min_final_score_for_entry", "value": cfg["min_final_score_for_entry"]},
                              *[{"key": f"model.{r.agent}", "value": f"{r.provider}/{r.model}"} for r in (stocks[0].agent_results.values() if stocks else []) if r.provider],
                              {"key": "place_order", "value": "false (research only)"}])
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xw:
        for name, df in (("Rankings", rank_df), ("Agent Scores", agent_scores_frame(stocks)), ("Events", events),
                         ("Data Quality", data_quality_frame(stocks)), ("Configuration", config_df)):
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) if len(df) else len(str(col))
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    rank_df.to_csv(latest, index=False)  # baseline for next run's rank_change
    return {"csv": csv_path, "xlsx": xlsx_path}
